"""Builds the Electify mini-project report (DOCX).

    cd documentation/final-report/work && python3 build_report.py [output.docx]
"""
import os
import sys
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

from docx_engine import Report, run_font, table_borders, TEXT_W
import content_a as A
import content_b as B
import content_c as C
import content_d as D

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "render", "report_build.docx")


def line(R, text, size=12, bold=True, after=8, before=0, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = R.doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    run_font(p.add_run(text), size=size, bold=bold)
    return p


def cover(R):
    for i, t in enumerate(A.TITLE_LINES):
        line(R, t, 16, after=22 if i == len(A.TITLE_LINES) - 1 else 4, before=6 if i == 0 else 0)
    for t in ["A PROJECT REPORT SUBMITTED TO", "SRM INSTITUTE OF SCIENCE & TECHNOLOGY",
              "IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE", "AWARD OF THE DEGREE OF",
              "MASTER OF COMPUTER APPLICATIONS"]:
        line(R, t, 12, after=9)
    line(R, "BY", 12, after=14, before=4)
    names = R.doc.add_paragraph()
    names.alignment = WD_ALIGN_PARAGRAPH.LEFT
    names.paragraph_format.left_indent = Inches(1.95)
    names.paragraph_format.space_after = Pt(18)
    names.paragraph_format.line_spacing = 1.15
    for i, (n, reg) in enumerate(A.TEAM):
        r = names.add_run(f"{n} ({reg})")
        run_font(r, size=12, bold=True)
        if i < len(A.TEAM) - 1:
            r.add_break()
    line(R, "UNDER THE GUIDANCE OF", 12, after=4)
    line(R, "Dr. V. NISHA", 12, after=12)
    p = R.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture("assets/srm_logo.png", width=Inches(2.68))
    for t in ["DEPARTMENT OF COMPUTER APPLICATIONS", "FACULTY OF SCIENCE AND HUMANITIES", "SRM INSTITUTE OF SCIENCE & TECHNOLOGY"]:
        line(R, t, 12, after=9)
    line(R, "Kattankulathur – 603 203", 12, bold=False, after=6)
    line(R, "Chennai, Tamil Nadu", 12, bold=False, after=22)
    line(R, "OCTOBER – 2026", 12, after=0)


def borderless_table(R, rows, col_w=(3.3, 3.3), size=11, bold_first=True, line_sp=1.6):
    t = R.doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, val="nil", sz=0, color="FFFFFF")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.width = Inches(col_w[ci])
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = line_sp
            run_font(p.add_run(val), size=size, bold=(bold_first and ri == 0))
    return t


def certificate(R):
    R.centered_heading("BONAFIDE CERTIFICATE", page_break=False, after=22)
    p = R.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing = 1.8
    p.paragraph_format.space_after = Pt(10)
    run_font(p.add_run("This is to certify that the project report titled “"))
    run_font(p.add_run(A.TITLE), bold=True)
    run_font(p.add_run("” is a bonafide work carried out by"))
    names = R.doc.add_paragraph()
    names.alignment = WD_ALIGN_PARAGRAPH.CENTER
    names.paragraph_format.line_spacing = 1.5
    names.paragraph_format.space_after = Pt(10)
    for i, (n, reg) in enumerate(A.TEAM):
        r = names.add_run(f"{n} ({reg})")
        run_font(r, bold=True)
        if i < len(A.TEAM) - 1:
            r.add_break()
    q = R.doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q.paragraph_format.line_spacing = 1.8
    q.paragraph_format.space_after = Pt(24)
    run_font(q.add_run("under my supervision for the award of the Degree of Master of Computer Applications. To my knowledge the work reported herein is the original work done by the students."))
    for _ in range(2):
        R.doc.add_paragraph().paragraph_format.space_after = Pt(0)
    borderless_table(R, [
        ["Dr. V. Nisha", "Dr. P. J. Arul Leena Rose"],
        ["Assistant Professor", "Professor & Head"],
        ["Department of Computer Applications", "Department of Computer Applications"],
        ["(GUIDE)", ""],
    ])
    for _ in range(5):
        R.doc.add_paragraph().paragraph_format.space_after = Pt(0)
    borderless_table(R, [["INTERNAL EXAMINER", "EXTERNAL EXAMINER"]])


def acknowledgement(R):
    R.centered_heading("ACKNOWLEDGEMENT", page_break=True, after=20)
    for t in A.ACKNOWLEDGEMENT:
        R.para(t)
    p = R.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.line_spacing = 1.3
    for i, (n, _) in enumerate(A.TEAM):
        r = p.add_run(n)
        run_font(r, bold=True)
        if i < len(A.TEAM) - 1:
            r.add_break()


def front_lists(R):
    R.centered_heading("TABLE OF CONTENTS", page_break=True, after=14)
    R.toc('TOC \\o "1-3" \\h \\z \\u')
    R.centered_heading("ABSTRACT", page_break=True, after=20)
    for t in A.ABSTRACT:
        R.para(t)
    R.para(f"**Keywords:** {A.KEYWORDS}", indent=0, before=6)
    R.centered_heading("LIST OF TABLES", page_break=True, after=14)
    R.toc('TOC \\h \\z \\t "Table Caption,1"')
    R.centered_heading("LIST OF FIGURES", page_break=True, after=14)
    R.toc('TOC \\h \\z \\t "Figure Caption,1"')


def build():
    R = Report()
    cover(R)
    R.new_section("lowerRoman", 1)          # certificate onwards: i, ii, iii …
    certificate(R)
    acknowledgement(R)
    front_lists(R)
    s = R.new_section("decimal", 1)        # chapters: 1, 2, 3 …
    # The first chapter heading must not add a second page break after the section break.
    A.chapter1(R)
    first_h1 = [p for p in R.doc.paragraphs if p.style.name == "Heading 1"][0]
    first_h1.paragraph_format.page_break_before = False
    A.chapter2(R); A.chapter3(R)
    B.chapter4(R); B.chapter5(R)
    C.chapter6(R); C.chapter7(R)
    D.chapter8(R); D.chapter9(R); D.chapter10(R); D.chapter11(R)
    # The cover section has no page number.
    R.doc.sections[0].footer.is_linked_to_previous = False
    R.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
