"""Small rendering engine: turns content blocks into a Word document that
follows the formatting of the department's reference report
(Times New Roman, 14/13/12 pt bold headings, justified body text with a
0.5 inch first-line indent, navy table headers, captioned figures)."""
import copy
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Courier New"
NAVY = "1F3864"
PAGE_W, PAGE_H = 8.27, 11.69
MARGINS = dict(left=1.0, right=0.66, top=1.0, bottom=1.01)
TEXT_W = PAGE_W - MARGINS["left"] - MARGINS["right"]


# ───────────────────────────────────────────────────────────── low-level XML
def _set_fonts(rpr_parent, name=FONT):
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, "get_or_add_rPr") else rpr_parent
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in list(fonts.attrib):
        del fonts.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), name)


def style_font(style, size=None, bold=None, italic=None, color="000000", name=FONT):
    style.font.name = name
    _set_fonts(style.element.get_or_add_rPr(), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color)


def run_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    _set_fonts(run._r.get_or_add_rPr(), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def table_borders(table, val="single", sz=4, color="808080"):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    existing = tblpr.find(qn("w:tblBorders"))
    if existing is not None:
        tblpr.remove(existing)
    tblpr.append(b)


def row_flags(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        h = OxmlElement("w:tblHeader")
        trpr.append(h)


def add_field(paragraph, instr, placeholder=""):
    r1 = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    r1._r.append(f1)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
    r2._r.append(it)
    r3 = paragraph.add_run()
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    r3._r.append(f2)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    r5._r.append(f3)
    for r in (r1, r2, r3, r4, r5):
        run_font(r)
    return paragraph


def page_number_format(section, fmt, start=1):
    sectpr = section._sectPr
    pg = sectpr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectpr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


def keep_next(p, on=True):
    p.paragraph_format.keep_with_next = on


# ───────────────────────────────────────────────────────────── document
class Report:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
        for k, v in MARGINS.items():
            setattr(s, f"{k}_margin", Inches(v))
        s.header_distance = Inches(0.5)
        s.footer_distance = Inches(0.45)
        self.numbering_part = self.doc.part.numbering_part
        self._list_abstract = {}

    # styles ----------------------------------------------------------------
    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        style_font(normal, 12)
        pf = normal.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(12)
        pf.line_spacing = 1.05
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Body text: justified with first-line indent, as in the reference
        body = st.add_style("Body Text Report", WD_STYLE_TYPE.PARAGRAPH)
        body.base_style = normal
        body.quick_style = True
        body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body.paragraph_format.first_line_indent = Inches(0.5)
        body.paragraph_format.widow_control = True
        for name, size, before, after in (("Heading 1", 14, 0, 14), ("Heading 2", 13, 10, 8), ("Heading 3", 12, 8, 6)):
            h = st[name]
            style_font(h, size, bold=True, italic=False)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.05
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.first_line_indent = Inches(0)
        st["Heading 1"].paragraph_format.page_break_before = True
        for name, align, bold in (("Figure Caption", WD_ALIGN_PARAGRAPH.CENTER, False), ("Table Caption", WD_ALIGN_PARAGRAPH.LEFT, False)):
            c = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            c.base_style = normal
            c.quick_style = True
            style_font(c, 12, bold=bold)
            c.paragraph_format.alignment = align
            c.paragraph_format.space_before = Pt(4 if name == "Figure Caption" else 6)
            c.paragraph_format.space_after = Pt(12 if name == "Figure Caption" else 4)
        tc = st.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        tc.base_style = normal
        style_font(tc, 11)
        tc.paragraph_format.space_after = Pt(0)
        tc.paragraph_format.line_spacing = 1.0
        code = st.add_style("Code Text", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        style_font(code, 8.5, name=MONO)
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.line_spacing = 1.0
        for i, (indent, bold) in enumerate(((0, True), (0.25, False), (0.5, False)), start=1):
            t = st.add_style(f"TOC {i}", WD_STYLE_TYPE.PARAGRAPH)
            t.element.set(qn("w:styleId"), f"TOC{i}")
            t.element.find(qn("w:name")).set(qn("w:val"), f"toc {i}")
            if qn("w:customStyle") in t.element.attrib:
                del t.element.attrib[qn("w:customStyle")]
            t.base_style = normal
            style_font(t, 12, bold=bold)
            t.paragraph_format.left_indent = Inches(indent)
            t.paragraph_format.space_after = Pt(5 if i == 1 else 3)
            t.paragraph_format.space_before = Pt(4 if i == 1 else 0)
            t.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W - 0.02), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        for name in ("List Bullet", "List Number"):
            l = st[name]
            style_font(l, 12)
            l.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            l.paragraph_format.space_after = Pt(3)
            l.paragraph_format.line_spacing = 1.05
        fs = st["Footer"]
        style_font(fs, 11)

    # sections --------------------------------------------------------------
    def new_section(self, page_fmt=None, start=1, footer=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.footer.is_linked_to_previous = False
        s.header.is_linked_to_previous = False
        fp = s.footer.paragraphs[0]
        for r in list(fp.runs):
            r._r.getparent().remove(r._r)
        if footer:
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(fp, "PAGE", "1")
        if page_fmt:
            page_number_format(s, page_fmt, start)
        return s

    # inline markup: **bold**, *italic*, `code` ------------------------------
    TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")

    def rich(self, p, text, size=None, base_bold=None):
        for part in self.TOKEN.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run_font(p.add_run(part[2:-2]), size=size, bold=True)
            elif part.startswith("`") and part.endswith("`"):
                run_font(p.add_run(part[1:-1]), name=MONO, size=(size or 12) - 1.5)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run_font(p.add_run(part[1:-1]), size=size, italic=True, bold=base_bold)
            else:
                run_font(p.add_run(part), size=size, bold=base_bold)
        return p

    # blocks ------------------------------------------------------------------
    def para(self, text, style="Body Text Report", align=None, size=None, bold=None, indent=None, after=None, before=None, keep=False):
        p = self.doc.add_paragraph(style=style)
        self.rich(p, text, size=size, base_bold=bold)
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.first_line_indent = Inches(indent)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if keep:
            keep_next(p)
        return p

    def heading(self, text, level):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        run_font(p.add_run(text), size={1: 14, 2: 13, 3: 12}[level], bold=True)
        return p

    def centered_heading(self, text, page_break=True, size=14, after=16):
        """Front-matter heading (not part of the TOC outline)."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.page_break_before = page_break
        run_font(p.add_run(text), size=size, bold=True)
        return p

    def _new_num(self, kind):
        """Fresh numbering instance so each list restarts at 1."""
        style = self.doc.styles["List Number" if kind == "number" else "List Bullet"]
        numpr = style.element.pPr.find(qn("w:numPr"))
        base_num_id = int(numpr.find(qn("w:numId")).get(qn("w:val")))
        numbering = self.numbering_part.element
        base = [n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == base_num_id][0]
        abs_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abs_id); num.append(a)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
        numbering.append(num)
        return new_id

    def items(self, entries, kind="bullet", after=10):
        style = "List Number" if kind == "number" else "List Bullet"
        num_id = self._new_num(kind) if kind == "number" else None
        ps = []
        for e in entries:
            p = self.doc.add_paragraph(style=style)
            self.rich(p, e)
            if num_id:
                ppr = p._p.get_or_add_pPr()
                numpr = OxmlElement("w:numPr")
                il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numpr.append(il)
                ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id)); numpr.append(ni)
                ppr.insert(0, numpr) if ppr.find(qn("w:pStyle")) is None else ppr.find(qn("w:pStyle")).addnext(numpr)
            ps.append(p)
        if ps:
            ps[-1].paragraph_format.space_after = Pt(after)
        return ps

    def table(self, caption, headers, rows, widths=None, size=11, align_cols=None, bold_first=False):
        cap = self.doc.add_paragraph(style="Table Caption")
        label, _, title = caption.partition(" | ")
        run_font(cap.add_run(label + " "), bold=True)
        run_font(cap.add_run(title), bold=True)
        keep_next(cap)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t)
        cell_margins(t)
        widths = widths or [TEXT_W / len(headers)] * len(headers)
        scale = TEXT_W / sum(widths)
        widths = [w * scale for w in widths]
        for ci, h in enumerate(headers):
            c = t.rows[0].cells[ci]
            shade(c, NAVY)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.style = self.doc.styles["Table Text"]
            run_font(p.add_run(h), size=size, bold=True, color="FFFFFF")
            keep_next(p)
        row_flags(t.rows[0], header=True)
        for ri, row in enumerate(rows, start=1):
            row_flags(t.rows[ri])
            for ci, val in enumerate(row):
                c = t.rows[ri].cells[ci]
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                lines = val.split("\n") if isinstance(val, str) else [str(val)]
                for li, line in enumerate(lines):
                    p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                    p.style = self.doc.styles["Table Text"]
                    if align_cols and ci in align_cols:
                        p.alignment = align_cols[ci]
                    self.rich(p, line, size=size, base_bold=True if (bold_first and ci == 0) else None)
        for row in t.rows:
            for ci, c in enumerate(row.cells):
                c.width = Inches(widths[ci])
        # Fixed layout so Word honours the widths
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(layout)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing = 0.8
        return t

    def figure(self, path, caption, width=6.3):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        keep_next(p)
        p.add_run().add_picture(path, width=Inches(min(width, TEXT_W)))
        cap = self.doc.add_paragraph(style="Figure Caption")
        label, _, title = caption.partition(" | ")
        run_font(cap.add_run(label + " "), bold=True)
        run_font(cap.add_run(title))
        return cap

    def code(self, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(t, sz=4, color="BFBFBF")
        cell_margins(t, 80, 80, 120, 120)
        c = t.rows[0].cells[0]
        shade(c, "F5F5F5")
        c.width = Inches(TEXT_W)
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph()
            p.style = self.doc.styles["Code Text"]
            run_font(p.add_run(line if line else " "), name=MONO, size=8.5)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        return t

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)

    def toc(self, instr):
        p = self.doc.add_paragraph()
        add_field(p, instr, "Right-click and choose Update Field to build this list.")
        return p

    def save(self, path):
        self.doc.core_properties.title = "Electify Mini Project Report"
        self.doc.core_properties.author = "Matheshwaran R, Suhas C, Nithyashree K, Devananda E"
        self.doc.core_properties.comments = ""
        self.doc.save(path)
