"""Builds the Electify MCA project presentation (16 slides, 16:9, editable).

    python3 artifacts/presentation/build_deck.py

Sources: the project report (documentation/final-report), the repository at
commit bfd9bd8, and the test evidence of 15 October 2026. Screenshots show
fictional demonstration data only.
"""
import copy
import os
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, "assets")
ICON = os.path.join(A, "icons")
OUT = os.path.join(HERE, "ELECTIFY_A_REAL-TIME_ELECTIVE_REGISTRATION_SYSTEM_WITH_TRANSACTIONAL_SEAT_MANAGEMENT_Project_Presentation.pptx")

# ── verified identity (from the team's report "Electify Final Report (1).docx", title confirmed by the student)
TITLE = "ELECTIFY: A REAL - TIME ELECTIVE REGISTRATION SYSTEM WITH TRANSACTIONAL SEAT MANAGEMENT"
STUDENT = "MATHESHWARAN R"
REGNO = "RA2532241010311"
GUIDE = "Dr. V. NISHA, M.Sc., M.Phil., Ph.D."
GUIDE_ROLE = "Assistant Professor, Department of Computer Applications"
DEPT = "Department of Computer Applications · Faculty of Science and Humanities"
INST = "SRM Institute of Science and Technology, Kattankulathur – 603 203"
COMPLETION_LETTER = os.environ.get("COMPLETION_LETTER")  # path to the official scan, when provided

# ── design system: Electify brand (globals.css) — Midnight dominant, Scholar Gold accent
INK = RGBColor(0x0E, 0x13, 0x30)       # Midnight
INK2 = RGBColor(0x1B, 0x22, 0x4A)      # raised midnight (cards on dark)
GOLD = RGBColor(0xE9, 0xB4, 0x4C)      # Scholar Gold
GOLD_SOFT = RGBColor(0xFB, 0xF1, 0xDC)
SLATE = RGBColor(0x55, 0x5B, 0x73)     # body muted
MIST = RGBColor(0xF3, 0xF4, 0xF8)      # card tint
LINE = RGBColor(0xD9, 0xDC, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DANGER = RGBColor(0xB4, 0x41, 0x3A)
OK = RGBColor(0x2F, 0x7D, 0x4F)
MUTED_ON_DARK = RGBColor(0xB9, 0xBE, 0xD3)
HEAD = "Georgia"
BODY = "Calibri"
MONO = "Consolas"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# ───────────────────────────────────────────────────────── primitives
def rgb_fill(shape, color):
    if color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color


def rect(s, x, y, w, h, fill=None, line=None, radius=None, lw=0.75, name=None, shadow=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = radius
    rgb_fill(shp, fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    if not shadow:
        _no_shadow(shp)
    else:
        _soft_shadow(shp)
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def _no_shadow(shp):
    sppr = shp._element.spPr
    eff = etree.SubElement(sppr, qn("a:effectLst"))


def _soft_shadow(shp):
    sppr = shp._element.spPr
    eff = etree.SubElement(sppr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="190500", dist="38100", dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="0E1330")
    etree.SubElement(clr, qn("a:alpha"), val="12000")


def oval(s, x, y, d, fill, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    rgb_fill(shp, fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    _no_shadow(shp)
    return shp


def text(s, x, y, w, h, content, size=18, color=INK, font=BODY, bold=False, italic=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_spacing=None, space_after=0, name=None, margin=0):
    """content: str | list of paragraphs; a paragraph is str or list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    if name:
        tb.name = name
    return tb


def bullets(s, x, y, w, h, items, size=18, color=INK, gap=8, bullet_color=GOLD, char="•"):
    """Real bulleted paragraphs (buChar) with hanging indent."""
    tb = text(s, x, y, w, h, [""], size=size, color=color)
    tf = tb.text_frame
    tf.paragraphs[0]._p.getparent().remove(tf.paragraphs[0]._p)
    for it in items:
        p = tf.add_paragraph()
        p.space_after = Pt(gap)
        runs = it if isinstance(it, list) else [(it, {})]
        for t, o in runs:
            r = p.add_run(); r.text = t
            r.font.name = o.get("font", BODY); r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False); r.font.color.rgb = o.get("color", color)
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(Emu(Inches(0.28)))); ppr.set("indent", str(-Emu(Inches(0.28))))
        bc = etree.SubElement(ppr, qn("a:buClr")); etree.SubElement(bc, qn("a:srgbClr"), val=str(bullet_color))
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(ppr, qn("a:buChar"), char=char)
    return tb


def icon(s, name, x, y, size, color="ink"):
    return s.shapes.add_picture(os.path.join(ICON, f"{name}-{color}.png"), Inches(x), Inches(y), Inches(size), Inches(size))


def icon_badge(s, name, x, y, d=0.62, bg=GOLD_SOFT, color="ink"):
    oval(s, x, y, d, bg)
    pad = d * 0.24
    icon(s, name, x + pad, y + pad, d - 2 * pad, color)


def arrow(s, x1, y1, x2, y2, color=SLATE, w=1.5, dash=False, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash"), val="dash")
    if head:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def picture_fit(s, path, x, y, w, h, crop_to=None):
    """Places an image inside (w,h) keeping aspect; returns the picture."""
    im = Image.open(path)
    iw, ih = im.size
    r = min(w / (iw / 100), h / (ih / 100))
    pw, ph = iw / 100 * r, ih / 100 * r
    return s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def screen(s, name, x, y, w, crop_h_ratio=None, border=True):
    """Screenshot in a light device frame; optional crop to a height/width ratio from the top."""
    src = os.path.join(A, "screens", name + ".png")
    opt = os.path.join(A, "screens_opt", name + ".jpg")
    os.makedirs(os.path.dirname(opt), exist_ok=True)
    im = Image.open(src).convert("RGB")
    if crop_h_ratio:
        im = im.crop((0, 0, im.width, int(im.width * crop_h_ratio)))
    im = im.resize((1920, int(1920 * im.height / im.width)), Image.LANCZOS)
    im.save(opt, quality=90, optimize=True)
    h = w * im.height / im.width
    if border:
        rect(s, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=WHITE, line=LINE, radius=0.03, shadow=True)
    pic = s.shapes.add_picture(opt, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.5)
    return h


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ───────────────────────────────────────────────────────── slide frames
def content_slide(n, kicker, title):
    s = prs.slides.add_slide(BLANK)
    rgb_fill(s.background, WHITE) if False else None
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    oval(s, 0.6, 0.47, 0.13, GOLD)
    text(s, 0.82, 0.39, 8.0, 0.3, kicker.upper(), size=12, color=SLATE, bold=True, name="Kicker")
    text(s, 0.6, 0.72, 11.6, 0.75, title, size=30, color=INK, font=HEAD, name="Title")
    s.shapes.add_picture(os.path.join(A, "electify-logo.png"), Inches(0.6), Inches(7.0), Inches(0.98), Inches(0.98 * 588 / 2600))
    text(s, 1.72, 6.99, 6.0, 0.25, "MCA Project Review · SRM Institute of Science and Technology", size=10, color=SLATE)
    text(s, 11.53, 6.99, 1.2, 0.25, f"{n:02d} / 16", size=11, color=SLATE, align=PP_ALIGN.RIGHT, name="Slide number")
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
    return s


def seat_grid(s, x, y, cols, rows, cell, gap, filled, gold_at, dim=INK2, on=WHITE):
    """The Electify motif: a block of seats, some taken, one gold."""
    k = 0
    for r in range(rows):
        for c in range(cols):
            col = GOLD if (r, c) == gold_at else (on if k in filled else dim)
            rect(s, x + c * (cell + gap), y + r * (cell + gap), cell, cell, fill=col, radius=0.18)
            k += 1


# ═════════════════════════════════════════════════════════ 1 · TITLE
s = dark_slide()
rect(s, 0.6, 0.55, 2.25, 0.82, fill=WHITE, radius=0.12)
picture_fit(s, os.path.join(A, "srm_logo.png"), 0.72, 0.62, 2.0, 0.68)
s.shapes.add_picture(os.path.join(A, "electify-logo-reverse.png"), Inches(3.15), Inches(0.78), Inches(1.6), Inches(1.6 * 588 / 2600))
text(s, 0.6, 1.75, 8.2, 0.3, "MCA PROJECT REVIEW", size=13, color=GOLD, bold=True)
text(s, 0.6, 2.15, 8.3, 2.3, TITLE, size=32, color=WHITE, font=HEAD, line_spacing=1.08, anchor=MSO_ANCHOR.TOP, name="Project title")
# identity block
text(s, 0.6, 4.75, 3.9, 0.28, "PRESENTED BY", size=11, color=MUTED_ON_DARK, bold=True)
text(s, 0.6, 5.05, 3.9, 0.4, STUDENT, size=20, color=WHITE, bold=True)
text(s, 0.6, 5.45, 3.9, 0.32, f"Register No. {REGNO}", size=15, color=MUTED_ON_DARK)
text(s, 4.7, 4.75, 4.6, 0.28, "UNDER THE GUIDANCE OF", size=11, color=MUTED_ON_DARK, bold=True)
text(s, 4.7, 5.05, 4.6, 0.4, GUIDE, size=18, color=WHITE, bold=True)
text(s, 4.7, 5.45, 4.6, 0.6, GUIDE_ROLE, size=15, color=MUTED_ON_DARK)
text(s, 0.6, 6.45, 9.0, 0.6, [DEPT, INST], size=12, color=MUTED_ON_DARK, line_spacing=1.1)
seat_grid(s, 9.75, 1.55, 5, 6, 0.46, 0.14, filled={0, 1, 2, 4, 5, 7, 8, 9, 11, 12, 13, 15, 16, 18, 19, 21, 23, 24, 26}, gold_at=(3, 2))
text(s, 9.75, 5.25, 2.9, 0.6, "One seat, one student — confirmed in a single database transaction.", size=12, color=MUTED_ON_DARK, italic=True)
notes(s, "Introduce the project title exactly as approved, yourself, and the guide. Electify is a web application for running elective registration in a department: it enforces seat limits exactly, keeps cohorts and academic years separate, and gives each staff role only the access it needs.")

# ═════════════════════════════════════════════════════════ 2 · ABSTRACT
s = content_slide(2, "Abstract", "Fair, exact elective registration for a whole department")
cols = [
    ("TriangleAlert", "The problem", "Seats are limited and the window is short. Forms and spreadsheets cannot enforce a seat limit at the moment of choice, and several cohorts run at once."),
    ("Workflow", "The solution", "A web application in which every registration, change and cancellation is one PostgreSQL transaction that locks the subject rows and applies capacity and class quotas."),
    ("CircleCheck", "Verified outcome", "136/136 server-side tests and 36/36 browser tests pass; 0 layout overflows at 7 screen sizes; 0 automated accessibility violations (runs of 15 Oct 2026)."),
]
for i, (ic, head, body) in enumerate(cols):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.8, 3.35, fill=MIST, radius=0.06)
    icon_badge(s, ic, x + 0.35, 2.05)
    text(s, x + 0.35, 2.85, 3.2, 0.4, head, size=20, bold=True)
    text(s, x + 0.35, 3.35, 3.15, 2.0, body, size=16, color=SLATE, line_spacing=1.08)
rect(s, 0.6, 5.35, 12.1, 0.85, fill=INK, radius=0.08)
text(s, 0.95, 5.45, 2.2, 0.65, "Objectives", size=15, color=GOLD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 2.55, 5.45, 10.0, 0.65, "Exact seat limits  ·  cohort and Year Server isolation  ·  scoped staff roles  ·  auditable records",
     size=15, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.4, 12.1, 0.28, "Built with  Next.js 16 · React 19 · TypeScript · Drizzle ORM · PostgreSQL · optional Supabase Realtime", size=11, color=SLATE)
notes(s, "Summarise the abstract: problem, solution, outcome. The numbers are from the test runs of 15 October 2026 on commit bfd9bd8: 136 server-side tests in 18 suites and 36 browser tests, all passing.")

# ═════════════════════════════════════════════════════════ 3 · INTRODUCTION
s = content_slide(3, "Introduction", "What Electify is and who uses it")
text(s, 0.6, 1.75, 5.6, 2.6, [
    "Electives let students shape part of their programme. Each semester, every class picks a subject from baskets such as Professional Elective I and Open Elective I, and each subject has a fixed number of seats.",
    "Electify runs that exercise on the department's own web application: the window, the seats, the choices and the records.",
], size=16, color=SLATE, line_spacing=1.08, space_after=10)
roles = [("Building2", "System administrator", "Year Servers, departments, staff accounts, audit log"),
         ("UserCog", "Course coordinator", "Events, baskets, subjects, class quotas, reports"),
         ("Users", "Class tutor", "Class roster, signup codes, live window, password resets"),
         ("GraduationCap", "Student", "Explores subjects, registers, changes or cancels, receipt")]
for i, (ic, r, d) in enumerate(roles):
    y = 1.75 + i * 1.18
    rect(s, 6.75, y, 5.95, 1.0, fill=MIST, radius=0.1)
    icon_badge(s, ic, 6.95, y + 0.19)
    text(s, 7.8, y + 0.14, 4.7, 0.36, r, size=17, bold=True)
    text(s, 7.8, y + 0.52, 4.75, 0.36, d, size=14, color=SLATE)
rect(s, 0.6, 4.35, 5.6, 2.35, fill=GOLD_SOFT, radius=0.06)
text(s, 0.9, 4.55, 5.0, 0.35, "Why it matters", size=17, bold=True)
bullets(s, 0.9, 4.98, 5.1, 1.65, ["Most students act in the first minutes; popular subjects fill fast",
                                  "A seat promised twice means one student is moved later",
                                  "Two cohorts can share a class name such as MCA F"], size=15, gap=4)
notes(s, "Domain: elective registration in a university department. Four roles use the system. The problem matters because of short windows, scarce seats and several cohorts running at once.")

# ═════════════════════════════════════════════════════════ 4 · PROBLEM STATEMENT
s = content_slide(4, "Problem statement", "Read, check, write: how a simple form overbooks a subject")
# lost-update illustration
text(s, 0.6, 1.7, 6.4, 0.35, "Two students submit at the same moment", size=16, bold=True)
for i, who in enumerate(["Student A", "Student B"]):
    y = 2.2 + i * 1.55
    rect(s, 0.6, y, 1.55, 1.15, fill=INK, radius=0.1)
    text(s, 0.6, y, 1.55, 1.15, who, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for j, (lab, col) in enumerate([("reads\n1 seat left", MIST), ("check\npasses", MIST), ("writes\n0 seats", MIST)]):
        bx = 2.5 + j * 1.5
        rect(s, bx, y + 0.12, 1.25, 0.9, fill=col, line=LINE, radius=0.1)
        text(s, bx, y + 0.12, 1.25, 0.9, lab, size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        arrow(s, bx - 0.3 if j else 2.18, y + 0.57, bx - 0.03, y + 0.57)
rect(s, 0.6, 5.35, 6.45, 0.8, fill=RGBColor(0xFC, 0xED, 0xEC), radius=0.1)
icon(s, "CircleX", 0.82, 5.53, 0.44, "ink")
text(s, 1.42, 5.35, 5.5, 0.8, [[("Both are confirmed. ", {"bold": True, "color": DANGER}), ("One seat, two students: the classic lost update.", {})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
text(s, 7.55, 1.7, 5.1, 0.35, "Gaps in the current process", size=16, bold=True)
bullets(s, 7.55, 2.15, 5.15, 3.4, [
    "Seat limits are not enforced when the choice is made",
    "Allocation is reconstructed later by staff, by hand",
    "Students cannot see remaining seats; staff cannot see progress",
    "Batches with the same class name can be mixed",
    "A browser value (such as a class cookie) must never widen access",
], size=16, gap=9)
rect(s, 7.55, 5.35, 5.15, 0.8, fill=INK, radius=0.1)
text(s, 7.8, 5.35, 4.8, 0.8, "Needed: enforce capacity at commit time, isolate cohorts, scope every staff action",
     size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Problem statement from Chapter 1.2 of the report. The lost-update anomaly: if an application reads a seat count, checks it and writes a new value in separate steps, two simultaneous requests can both succeed. Electify must also keep cohorts apart and must not trust values sent by the browser.")

# ═════════════════════════════════════════════════════════ 5 · OBJECTIVES AND SCOPE
s = content_slide(5, "Objectives and scope", "What the project set out to deliver")
objs = [("Armchair", "Never exceed a subject's capacity or a class quota, even under simultaneous submissions"),
        ("Layers", "Model faculties → programmes → cohorts → classes; one Year Server per academic year"),
        ("Ticket", "Onboard students only through a tutor-issued class signup code"),
        ("Timer", "Enforce the registration window on the server clock"),
        ("ScrollText", "Issue a numbered receipt; keep an audit log and a registration replay"),
        ("ShieldCheck", "Limit every staff action to the classes the role and assignments allow"),
        ("TestTubeDiagonal", "Verify with automated tests against a real PostgreSQL database")]
text(s, 0.6, 1.65, 7.6, 0.35, "Core objectives — all implemented and verified", size=16, bold=True)
for i, (ic, t) in enumerate(objs):
    y = 2.1 + i * 0.66
    icon(s, ic, 0.62, y + 0.06, 0.36)
    text(s, 1.15, y, 6.95, 0.6, t, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.55, 1.65, 4.15, 2.45, fill=MIST, radius=0.06)
text(s, 8.85, 1.85, 3.6, 0.35, "In scope", size=16, bold=True)
bullets(s, 8.85, 2.25, 3.65, 1.8, ["Four roles, one department workflow", "Year Servers and cohort isolation", "Live monitoring, reports, replay", "Subject explorer with seat counts"], size=14, gap=3)
rect(s, 8.55, 4.3, 4.15, 2.35, fill=GOLD_SOFT, radius=0.06)
text(s, 8.85, 4.5, 3.6, 0.35, "Out of scope (future work)", size=16, bold=True)
bullets(s, 8.85, 4.9, 3.65, 1.7, ["Preference-ranked allocation", "Central student-system integration", "E-mail or SMS notifications", "Production deployment and load tests"], size=14, gap=3)
notes(s, "Objectives are from section 1.3.1 of the report and are verifiable: each maps to tests in Chapter 8. Allocation is first-come within the window; preference-based allocation, integration and notifications are explicitly future work. The system has not been deployed for a live registration cycle.")

# ═════════════════════════════════════════════════════════ 6 · EXISTING SYSTEM
s = content_slide(6, "Existing system and its limitations", "Today: forms, spreadsheets and manual allocation")
steps = [("ClipboardCheck", "Paper form or\nonline form"), ("FileCheck", "Shared\nspreadsheet"), ("Users", "Staff allocate\nby hand"), ("TriangleAlert", "Disputes and\nmanual moves")]
for i, (ic, lab) in enumerate(steps):
    x = 0.6 + i * 3.1
    rect(s, x, 1.85, 2.55, 1.55, fill=MIST, radius=0.1)
    icon(s, ic, x + 0.98, 2.0, 0.55)
    text(s, x, 2.62, 2.55, 0.7, lab, size=15, align=PP_ALIGN.CENTER, line_spacing=1.0)
    if i < 3:
        arrow(s, x + 2.6, 2.62, x + 3.05, 2.62, color=SLATE, w=1.75)
lims = [("No limit at the moment of choice", "A form records a request; whether a seat exists is decided later"),
        ("Opaque to students", "Nobody can see how many seats remain before choosing"),
        ("No live view for staff", "Who has registered is known only after the window closes"),
        ("Changes handled informally", "Withdrawals leave no reliable record of what changed and when"),
        ("Cohorts can mix", "Different batches end up in the same sheet or form"),
        ("All-or-nothing access", "Anyone who can edit the sheet can alter anyone's entry")]
for i, (h, d) in enumerate(lims):
    cx = 0.6 + (i % 3) * 4.1
    cy = 3.85 + (i // 3) * 1.45
    rect(s, cx, cy, 3.8, 1.25, fill=WHITE, line=LINE, radius=0.08)
    oval(s, cx + 0.25, cy + 0.28, 0.16, DANGER)
    text(s, cx + 0.55, cy + 0.17, 3.1, 0.35, h, size=15, bold=True)
    text(s, cx + 0.55, cy + 0.55, 3.1, 0.65, d, size=13, color=SLATE, line_spacing=1.0)
notes(s, "Section 3.1 of the report: where a department has no dedicated system, choices are collected on paper, in a shared spreadsheet, or through a general-purpose form. These work for small classes but fail when demand exceeds capacity, when cohorts share subjects and when students change their minds.")

# ═════════════════════════════════════════════════════════ 7 · PROPOSED SYSTEM
s = content_slide(7, "Proposed system and key features", "Enforce the rule where the data lives")
feats = [("Armchair", "Transactional seat allocation", "Rows locked in id order; capacity and class quotas checked inside one transaction"),
         ("Server", "Year Servers and cohort isolation", "Each academic year is a workspace; composite keys keep a class inside its cohort"),
         ("Ticket", "Class signup codes", "A code names one class; uses are claimed atomically, so the limit holds"),
         ("Search", "Subject explorer", "Students read about subjects and see live seat counts before the window opens"),
         ("Activity", "Live monitor and replay", "Turnout, seat saturation and an ordered replay of who took which seat"),
         ("ScrollText", "Receipts and audit log", "Unique receipt numbers, JSON snapshots and an audit trail of staff actions")]
for i, (ic, h, d) in enumerate(feats):
    cx = 0.6 + (i % 3) * 4.1
    cy = 1.7 + (i // 3) * 2.15
    rect(s, cx, cy, 3.8, 1.95, fill=MIST, radius=0.06)
    icon_badge(s, ic, cx + 0.3, cy + 0.28, d=0.58)
    text(s, cx + 1.05, cy + 0.3, 2.6, 0.55, h, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
    text(s, cx + 0.3, cy + 1.0, 3.3, 0.9, d, size=14, color=SLATE, line_spacing=1.03)
rect(s, 0.6, 6.12, 12.1, 0.62, fill=WHITE, line=LINE, radius=0.12)
text(s, 0.85, 6.12, 11.7, 0.62, [[("Status  ", {"bold": True, "color": OK}), ("All six features are implemented and covered by automated tests.   ", {}),
                                   ("Planned  ", {"bold": True, "color": SLATE}), ("preference-ranked allocation · notifications", {"color": SLATE})]],
     size=14, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Section 3.2 and Chapter 5–6 of the report. The difference from the existing process: the seat rule is enforced at commit time by the database, not reconstructed afterwards. Realtime seat updates use Supabase Realtime when it is configured; pages also refresh every 30 seconds, so correctness never depends on it.")

# ═════════════════════════════════════════════════════════ 8 · ARCHITECTURE
s = content_slide(8, "System architecture", "Five layers, one rule: the server decides")
cols = [("Browser", INK, ["Student portal", "Staff portal"], "HTTPS · httpOnly cookie"),
        ("Routing proxy", INK2, ["src/proxy.ts", "Verifies JWT", "Role route guards"], ""),
        ("Next.js 16 App Router", INK2, ["Server Components", "Server Actions", "Route handlers"], ""),
        ("Domain modules", INK2, ["auth · workspace", "student-events", "registration-window", "seat-ops"], ""),
        ("Data", INK, ["Drizzle ORM", "PostgreSQL", "Row locks · constraints"], "")]
cw, gap, top, ch = 2.15, 0.34, 1.75, 3.5
for i, (head, col, items, foot) in enumerate(cols):
    x = 0.6 + i * (cw + gap)
    rect(s, x, top, cw, ch, fill=MIST, radius=0.06)
    rect(s, x, top, cw, 0.62, fill=col, radius=0.18)
    text(s, x, top, cw, 0.62, head, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for j, it in enumerate(items):
        rect(s, x + 0.17, top + 0.85 + j * 0.62, cw - 0.34, 0.5, fill=WHITE, line=LINE, radius=0.2)
        text(s, x + 0.17, top + 0.85 + j * 0.62, cw - 0.34, 0.5, it, size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             font=MONO if ("." in it and "·" not in it) or it.startswith("src") else BODY)
    if foot:
        text(s, x, top + ch - 0.45, cw, 0.3, foot, size=11, color=SLATE, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + cw + 0.03, top + 1.6, x + cw + gap - 0.03, top + 1.6, color=INK, w=1.75)
# realtime side channel
ry = 5.75
rect(s, 4.85, ry, 5.6, 0.85, fill=GOLD_SOFT, line=GOLD, radius=0.12, lw=1)
icon(s, "RadioTower", 5.05, ry + 0.2, 0.45)
text(s, 5.65, ry, 4.7, 0.85, [[("Supabase Realtime (optional)  ", {"bold": True}), ("database changes pushed to open pages; pages also refresh every 30 s", {"color": SLATE, "size": 13})]],
     size=14, anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 11.55, top + ch + 0.03, 11.55, ry + 0.42, color=GOLD, dash=True, head=False)
arrow(s, 11.55, ry + 0.42, 10.48, ry + 0.42, color=GOLD, dash=True)
arrow(s, 4.82, ry + 0.42, 1.68, ry + 0.42, color=GOLD, dash=True, head=False)
arrow(s, 1.68, ry + 0.42, 1.68, top + ch + 0.03, color=GOLD, dash=True)
notes(s, "Every request passes the edge proxy (src/proxy.ts), which verifies the signed session cookie and routes by role. Pages are React Server Components; all changes go through Server Actions that call domain modules in src/lib. Drizzle ORM talks to PostgreSQL, which holds the row locks and constraints. Supabase Realtime is optional and only improves freshness.")

# ═════════════════════════════════════════════════════════ 9 · TECH STACK
s = content_slide(9, "Technology stack and tools", "Verified from package.json and the repository")
stack = [("Code2", "Language", ["TypeScript 5"]),
         ("Layers", "Framework", ["Next.js 16.2 (App Router)", "React 19.2"]),
         ("MonitorSmartphone", "Interface", ["Tailwind CSS 4 · Base UI", "Lucide icons · Framer Motion"]),
         ("FileCheck", "Forms and validation", ["React Hook Form 7", "Zod 4 (server-side)"]),
         ("Database", "Data", ["PostgreSQL", "Drizzle ORM 0.45 · drizzle-kit"]),
         ("KeyRound", "Authentication", ["jose — JWT (HS256)", "bcryptjs (cost 12)"]),
         ("RadioTower", "Live updates and documents", ["Supabase JS (Realtime)", "jsPDF — receipts and reports"]),
         ("TestTubeDiagonal", "Testing", ["node:test via tsx", "Playwright 1.64 · axe-core"]),
         ("Boxes", "Tooling", ["Git · npm · ESLint", "TypeScript compiler (tsc)"])]
for i, (ic, cat, items) in enumerate(stack):
    cx = 0.6 + (i % 3) * 4.1
    cy = 1.7 + (i // 3) * 1.68
    rect(s, cx, cy, 3.8, 1.48, fill=MIST, radius=0.07)
    icon(s, ic, cx + 0.28, cy + 0.27, 0.42)
    text(s, cx + 0.85, cy + 0.22, 2.8, 0.36, cat.upper(), size=12, color=SLATE, bold=True)
    text(s, cx + 0.85, cy + 0.58, 2.85, 0.85, items, size=15, line_spacing=1.05)
text(s, 0.6, 6.7, 12.1, 0.28, "Hosting: documented for a Node.js platform with pooled PostgreSQL; not yet deployed for a live registration cycle.", size=12, color=SLATE)
notes(s, "Versions are those declared in package.json. Deployment is documented (docs/deployment.md) but the system has only been run in development environments.")

# ═════════════════════════════════════════════════════════ 10 · MODULES AND ROLES
s = content_slide(10, "System modules and user roles", "Each role sees only its own part of the department")
roles = [("Building2", "Administrator", "Every class", ["Year Servers and cohorts", "Faculties and departments", "Users and invite codes", "Audit log, maintenance mode"]),
         ("UserCog", "Coordinator", "Own programme", ["Registration events", "Baskets and subjects", "Class seat quotas", "Templates, reports, monitor"]),
         ("Users", "Class tutor", "Assigned classes", ["Class roster and CSV import", "Signup codes, password resets", "Live window controls", "Class reports"]),
         ("GraduationCap", "Student", "Own registration", ["Join a class with a code", "Subject explorer", "Register, change, cancel", "Receipt (PDF or print)"])]
for i, (ic, r, scope, mods) in enumerate(roles):
    cx = 0.6 + i * 3.08
    rect(s, cx, 1.7, 2.85, 3.55, fill=MIST, radius=0.06)
    icon_badge(s, ic, cx + 0.25, 1.95, d=0.6)
    text(s, cx + 1.0, 1.95, 1.8, 0.36, r, size=17, bold=True)
    text(s, cx + 1.0, 2.3, 1.8, 0.3, scope, size=12, color=SLATE)
    bullets(s, cx + 0.25, 2.9, 2.5, 2.3, mods, size=16, gap=9)
rect(s, 0.6, 5.55, 12.1, 0.95, fill=INK, radius=0.08)
text(s, 0.9, 5.55, 2.4, 0.95, "Shared server core", size=15, color=GOLD, bold=True, anchor=MSO_ANCHOR.MIDDLE)
chips = ["Sessions", "Staff workspace", "Eligibility", "Window rule", "Seat operations", "Audit and replay"]
for j, c in enumerate(chips):
    x = 3.25 + j * 1.55
    rect(s, x, 5.79, 1.42, 0.47, fill=INK2, radius=0.4)
    text(s, x, 5.79, 1.42, 0.47, c, size=12, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Table 4.4 of the report. A tutor can manage classes from any programme of their department, up to a limit set by the administrator (ten by default). The root administrator account cannot be deleted, deactivated or demoted. All roles share the same server core in src/lib.")

# ═════════════════════════════════════════════════════════ 11 · WORKFLOW
s = content_slide(11, "Application workflow", "From sign-in to receipt: one registration")
def step(x, y, w, h, t, fill=MIST, color=INK, bold=False, size=13):
    rect(s, x, y, w, h, fill=fill, line=None if fill != WHITE else LINE, radius=0.14)
    text(s, x, y, w, h, t, size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
def decision(x, y, w, h, t):
    d = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(x), Inches(y), Inches(w), Inches(h))
    rgb_fill(d, GOLD_SOFT); d.line.color.rgb = GOLD; d.line.width = Pt(1); _no_shadow(d)
    text(s, x + 0.15, y, w - 0.3, h, t, size=12, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
text(s, 0.6, 1.6, 6, 0.3, "STUDENT", size=12, color=SLATE, bold=True)
r1 = 2.0
step(0.6, r1, 1.65, 0.9, "Sign in\nregister no. + password", fill=INK, color=WHITE, bold=True)
decision(2.55, r1 - 0.12, 1.9, 1.14, "Event for\nmy class?")
decision(4.75, r1 - 0.12, 1.9, 1.14, "Window open?\n(server clock)")
step(6.95, r1, 1.85, 0.9, "Pick one subject\nper basket")
step(9.1, r1, 1.6, 0.9, "Review and\nconfirm", fill=INK, color=WHITE, bold=True)
for a, b in [(2.25, 2.55), (4.45, 4.75), (6.65, 6.95), (8.8, 9.1)]:
    arrow(s, a + 0.02, r1 + 0.45, b - 0.02, r1 + 0.45, color=INK)
text(s, 2.55, r1 + 1.05, 1.9, 0.5, "no → class's own events only", size=11, color=SLATE, align=PP_ALIGN.CENTER)
text(s, 4.75, r1 + 1.05, 1.9, 0.5, "no → countdown or closed notice", size=11, color=SLATE, align=PP_ALIGN.CENTER)
# transaction lane
rect(s, 0.6, 3.75, 12.1, 2.95, fill=WHITE, line=LINE, radius=0.04)
text(s, 0.85, 3.88, 8, 0.3, "SERVER — ONE DATABASE TRANSACTION", size=12, color=SLATE, bold=True)
arrow(s, 9.9, r1 + 0.92, 9.9, 3.73, color=INK)
r2 = 4.45
step(0.85, r2, 2.0, 1.05, "Re-check eligibility\nand window")
step(3.15, r2, 2.15, 1.05, "Lock registration and\nsubjects (FOR UPDATE,\nin id order)")
decision(5.6, r2 - 0.1, 2.05, 1.25, "Seat and class\nquota available?")
step(7.95, r2, 2.05, 1.05, "Write choice, replay\nand audit rows")
step(10.3, r2, 2.15, 1.05, "Receipt number\n(advisory lock) →\nCOMMIT", fill=OK, color=WHITE, bold=True)
for a, b in [(2.85, 3.15), (5.3, 5.6), (7.65, 7.95), (10.0, 10.3)]:
    arrow(s, a + 0.02, r2 + 0.52, b - 0.02, r2 + 0.52, color=INK)
text(s, 7.6, r2 + 0.1, 0.4, 0.3, "yes", size=11, color=OK, bold=True)
arrow(s, 6.62, r2 + 1.16, 6.62, 5.95, color=DANGER)
step(5.25, 5.95, 2.75, 0.6, "No: ROLLBACK, code SEAT_TAKEN — nothing is written, picks are kept", fill=RGBColor(0xFC, 0xED, 0xEC), size=11)
step(10.3, 5.85, 2.15, 0.7, "Receipt shown;\nseats update", fill=GOLD_SOFT, bold=True)
arrow(s, 11.37, r2 + 1.07, 11.37, 5.83, color=INK)
notes(s, "registerElectives in src/app/actions/register.ts. Cheap checks happen first without locks. Inside the transaction the registration row and the chosen subjects are locked in a fixed order, so two transactions can never deadlock. takeSeat applies the subject capacity and any class quota. If any step fails the whole transaction rolls back. Changes take the new seat before releasing the old one; cancellations keep the record and return the seat once.")

# ═════════════════════════════════════════════════════════ 12 · DATABASE
s = content_slide(12, "Database design and data management", "Twelve core tables; integrity enforced by PostgreSQL")
grid = {
    (0, 0): ("academic_sessions", "Year Server: start / end year, status", "green"),
    (1, 0): ("registration_events", "FK cohort, session · window, status", "gold"),
    (2, 0): ("elective_groups", "FK event · min / max choices", "gold"),
    (3, 0): ("electives", "FK group · max / available seats", "gold"),
    (0, 1): ("cohort_session_terms", "FK cohort, session · year, semester", "green"),
    (1, 1): ("event_sections", "FK (event, cohort), (class, cohort)", "gold"),
    (2, 1): ("registrations", "FK student, event · receipt", "ink"),
    (3, 1): ("elective_seat_allocations", "FK elective, class · quota", "gold"),
    (0, 2): ("academic_batches", "FK programme · admission year", "green"),
    (1, 2): ("sections", "FK cohort · class label (F, G…)", "green"),
    (2, 2): ("users", "role · register no. · FK class", "ink"),
    (3, 2): ("student_registrations", "FK student, event, group, elective", "ink"),
}
tint = {"green": RGBColor(0xE6, 0xF2, 0xEA), "gold": GOLD_SOFT, "ink": RGBColor(0xE6, 0xE8, 0xF2)}
bx0, by0, bw, bh, gx, gy = 0.6, 1.7, 2.02, 1.2, 0.18, 0.5
def bpos(c, r):
    return bx0 + c * (bw + gx), by0 + r * (bh + gy)
for (c, r), (name, sub, t) in grid.items():
    x, y = bpos(c, r)
    rect(s, x, y, bw, bh, fill=WHITE, line=LINE, radius=0.08)
    rect(s, x, y, bw, 0.5, fill=tint[t], radius=0.18)
    text(s, x + 0.04, y, bw - 0.08, 0.5, name.replace("seat_allocations", "seat_\nallocations") if len(name) > 21 else name, size=11 if len(name) > 21 else 11.5, font=MONO, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.1, y + 0.55, bw - 0.2, 0.6, sub, size=11, color=SLATE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
def link(c1, r1, c2, r2):
    x1, y1 = bpos(c1, r1); x2, y2 = bpos(c2, r2)
    if r1 == r2:
        a = arrow(s, x1 + bw, y1 + bh / 2, x2, y2 + bh / 2, color=SLATE, w=1.25, head=False)
    else:
        a = arrow(s, x1 + bw / 2, y1 + bh, x2 + bw / 2, y2, color=SLATE, w=1.25, head=False)
for l in [((0, 0), (1, 0)), ((1, 0), (2, 0)), ((2, 0), (3, 0)), ((3, 0), (3, 1)), ((0, 0), (0, 1)), ((0, 1), (0, 2)),
          ((0, 2), (1, 2)), ((1, 0), (1, 1)), ((1, 1), (1, 2)), ((1, 2), (2, 2)), ((2, 1), (2, 2)), ((2, 2), (3, 2))]:
    link(*l[0], *l[1])
text(s, 0.6, 6.65, 8.6, 0.3, "Lines are foreign keys (one-to-many). Colours: academic structure · events and seats · people and registrations.", size=11, color=SLATE)
px = 9.85
rect(s, px, 1.7, 2.85, 4.9, fill=INK, radius=0.05)
text(s, px + 0.25, 1.9, 2.4, 0.35, "Integrity guarantees", size=16, color=GOLD, bold=True)
bullets(s, px + 0.25, 2.38, 2.45, 4.1, [
    [("available_seats ≥ 0 ", {"bold": True, "color": WHITE}), ("on subjects and quotas", {"color": MUTED_ON_DARK})],
    [("One choice ", {"bold": True, "color": WHITE}), ("per student, event and basket (unique)", {"color": MUTED_ON_DARK})],
    [("Unique ", {"bold": True, "color": WHITE}), ("receipt number per registration", {"color": MUTED_ON_DARK})],
    [("Composite keys ", {"bold": True, "color": WHITE}), ("keep a class, a student and a code inside one cohort", {"color": MUTED_ON_DARK})],
    [("Reviewed SQL migrations ", {"bold": True, "color": WHITE}), ("with rollback scripts", {"color": MUTED_ON_DARK})],
], size=13, gap=8, color=WHITE)
notes(s, "Simplified from src/lib/db/schema.ts (about 25 tables in total; login attempts, invite codes, audit logs, templates and subject profiles are omitted here). A cohort is a programme plus admission year; a class is a section of it; a Year Server links cohorts to an academic year with their study year and semester.")

# ═════════════════════════════════════════════════════════ 13 · SECURITY
s = content_slide(13, "Security, validation and reliability", "Protections implemented in the code")
sec = [("Lock", "Passwords", "bcrypt (cost 12); policy blocks the register number and common defaults"),
       ("KeyRound", "Sessions", "Signed JWT in an httpOnly cookie, 2 hours; one active session per user"),
       ("Hourglass", "Rate limiting", "5 failed sign-ins lock the identifier for 15 minutes"),
       ("ShieldCheck", "Scoped authorisation", "Role and class scope re-checked in every action; a cookie never grants access"),
       ("FileCheck", "Validation", "Zod schemas on input; eligibility re-checked inside the transaction"),
       ("Database", "Transactions", "Row locks in a fixed order; constraints as a final guard"),
       ("RotateCcw", "No double submits", "Already-registered and already-cancelled states are refused"),
       ("ScrollText", "Audit and replay", "Staff actions and every seat change are recorded")]
for i, (ic, h, d) in enumerate(sec):
    cx = 0.6 + (i % 4) * 3.08
    cy = 1.7 + (i // 4) * 2.2
    rect(s, cx, cy, 2.85, 2.0, fill=MIST, radius=0.07)
    icon_badge(s, ic, cx + 0.25, cy + 0.25, d=0.55)
    text(s, cx + 0.95, cy + 0.25, 1.85, 0.55, h, size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
    text(s, cx + 0.25, cy + 0.95, 2.4, 1.0, d, size=13, color=SLATE, line_spacing=1.02)
rect(s, 0.6, 6.12, 12.1, 0.62, fill=GOLD_SOFT, radius=0.12)
text(s, 0.85, 6.12, 11.7, 0.62, [[("Known gap  ", {"bold": True}), ("The 8-character, letter-and-number policy falls short of NIST SP 800-63B-4 (15 characters, no composition rules); future work.", {"color": SLATE})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Each item is implemented: lib/auth.ts, lib/password-policy.ts, lib/workspace.ts, actions/register.ts. A revoked or replaced session is rejected on the next request and the user is sent to sign in. Root-administrator protection prevents deleting, deactivating or demoting the root account. The NIST gap is stated in section 4.6 of the report.")

# ═════════════════════════════════════════════════════════ 14 · IMPLEMENTATION, TESTING, OUTCOMES
s = content_slide(14, "Implementation, testing and outcomes", "Verified on 15 October 2026 (commit bfd9bd8)")
h1 = screen(s, "13_electives_selected", 0.66, 1.75, 4.35, crop_h_ratio=0.48)
text(s, 0.6, 1.75 + h1 + 0.1, 4.5, 0.28, "Student picker: one subject chosen in each basket", size=11, color=SLATE)
y2 = 1.75 + h1 + 0.58
h2 = screen(s, "24_tutor_window", 0.66, y2, 4.35, crop_h_ratio=0.48)
text(s, 0.6, y2 + h2 + 0.1, 4.5, 0.28, "Tutor live window controlling the registration", size=11, color=SLATE)
stats = [("136 / 136", "server-side tests pass", "18 suites, real PostgreSQL"),
         ("36 / 36", "browser tests pass", "Playwright, every role, with reloads"),
         ("0", "layout overflows", "22 pages × 7 screen sizes"),
         ("0", "accessibility violations", "axe-core, WCAG 2.1 A/AA")]
for i, (big, lab, sub) in enumerate(stats):
    cx = 5.55 + (i % 2) * 3.6
    cy = 1.72 + (i // 2) * 1.62
    rect(s, cx, cy, 3.4, 1.45, fill=MIST, radius=0.07)
    text(s, cx + 0.28, cy + 0.15, 3.0, 0.65, big, size=32, font=HEAD, bold=True)
    text(s, cx + 0.28, cy + 0.78, 3.0, 0.3, lab, size=14, bold=True)
    text(s, cx + 0.28, cy + 1.07, 3.0, 0.3, sub, size=12, color=SLATE)
rect(s, 5.55, 5.05, 7.0, 1.65, fill=INK, radius=0.06)
text(s, 5.85, 5.2, 6.5, 0.32, "What testing found", size=15, color=GOLD, bold=True)
bullets(s, 5.85, 5.58, 6.5, 1.1, [
    [("14 defects found, 13 fixed ", {"bold": True, "color": WHITE}), ("— incl. a crashing countdown page and a session-revocation security fault", {"color": MUTED_ON_DARK})],
    [("Not yet tested: ", {"bold": True, "color": WHITE}), ("load, realtime with Supabase, other browsers, real users", {"color": MUTED_ON_DARK})],
], size=13, gap=4, color=WHITE)
notes(s, "All numbers come from the recorded runs (evidence in documentation/final-report/work/evidence). Screenshots show fictional demonstration data. Concurrency tests start several registrations at once for the last seat: exactly one succeeds. The one open defect (D7) is a reduced-motion hydration warning on the sign-in page.")

# ═════════════════════════════════════════════════════════ 15 · PROJECT COMPLETION LETTER
s = content_slide(15, "Official document", "PROJECT COMPLETION LETTER")
if COMPLETION_LETTER and os.path.exists(COMPLETION_LETTER):
    im = Image.open(COMPLETION_LETTER).convert("RGB")
    # crop only blank margins
    from PIL import ImageChops
    bg = Image.new("RGB", im.size, im.getpixel((0, 0)))
    bbox = ImageChops.difference(im, bg).getbbox()
    if bbox:
        pad = int(0.02 * max(im.size))
        im = im.crop((max(0, bbox[0] - pad), max(0, bbox[1] - pad), min(im.width, bbox[2] + pad), min(im.height, bbox[3] + pad)))
    letter = os.path.join(A, "completion_letter_cropped.png")
    im.save(letter)
    rect(s, 3.9, 1.6, 5.55, 5.25, fill=MIST, radius=0.03)
    picture_fit(s, letter, 4.0, 1.7, 5.35, 5.05)
else:
    ph = rect(s, 3.65, 1.65, 6.05, 5.1, fill=MIST, line=SLATE, radius=0.03, lw=1.25)
    ph.line.dash_style = 4  # dash
    icon(s, "FileCheck", 6.32, 2.65, 0.7)
    text(s, 3.9, 3.55, 5.55, 0.5, "Insert the official Project Completion Letter", size=18, bold=True, align=PP_ALIGN.CENTER)
    text(s, 3.9, 4.1, 5.55, 1.3, "Placeholder — the signed and sealed letter issued by the institution goes here, unchanged. It was not available when this deck was built.",
         size=14, color=SLATE, align=PP_ALIGN.CENTER, line_spacing=1.08)
notes(s, "Show the official Project Completion Letter. Do not retype or recreate it; insert the original scan.")

# ═════════════════════════════════════════════════════════ 16 · THANK YOU
s = dark_slide()
seat_grid(s, 9.75, 1.55, 5, 6, 0.46, 0.14, filled={0, 1, 2, 4, 5, 7, 8, 9, 11, 12, 13, 15, 16, 18, 19, 21, 23, 24, 26}, gold_at=(3, 2))
s.shapes.add_picture(os.path.join(A, "electify-logo-reverse.png"), Inches(0.6), Inches(0.78), Inches(1.6), Inches(1.6 * 588 / 2600))
text(s, 0.6, 2.25, 8.5, 1.4, "Thank You", size=66, color=WHITE, font=HEAD)
text(s, 0.6, 3.65, 8.5, 0.4, "Questions are welcome.", size=18, color=MUTED_ON_DARK, italic=True)
text(s, 0.6, 4.75, 3.9, 0.28, "STUDENT NAME", size=11, color=MUTED_ON_DARK, bold=True)
text(s, 0.6, 5.05, 3.9, 0.4, STUDENT, size=20, color=WHITE, bold=True)
text(s, 4.7, 4.75, 4.6, 0.28, "REGISTRATION NUMBER", size=11, color=MUTED_ON_DARK, bold=True)
text(s, 4.7, 5.05, 4.6, 0.4, REGNO, size=20, color=WHITE, bold=True)
text(s, 0.6, 6.45, 9.0, 0.6, [DEPT, INST], size=12, color=MUTED_ON_DARK, line_spacing=1.1)
notes(s, "Close and invite questions.")

prs.core_properties.title = TITLE
prs.core_properties.author = STUDENT
prs.core_properties.subject = "MCA project review presentation"
prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides))
